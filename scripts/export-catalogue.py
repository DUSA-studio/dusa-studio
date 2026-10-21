"""Export the Growth Engine catalogue (xlsx) to the JSON the Ride with DUSA builder reads.

Run after changing prices in DUSA-Growth-Engine-Catalogue.xlsx (and after the workbook has been
recalculated, i.e. opened and saved in Excel or run through LibreOffice):

    python3 scripts/export-catalogue.py ../DUSA-Growth-Engine-Catalogue.xlsx src/data/catalogue-au.json
"""
import json
import sys
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "../DUSA-Growth-Engine-Catalogue.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "src/data/catalogue-au.json"
wb = load_workbook(src, data_only=True)
c, i = wb["Catalogue"], wb["Inputs"]
items = {}
for r in range(5, c.max_row + 1):
    sid = c.cell(r, 1).value
    if not sid or "-" not in str(sid):
        continue
    g = c.cell(r, 8).value
    items[sid] = {
        "service": c.cell(r, 3).value, "size": c.cell(r, 4).value, "incl": c.cell(r, 5).value,
        "setup": c.cell(r, 6).value or 0, "monthly": c.cell(r, 7).value or 0,
        "guided": g if isinstance(g, (int, float)) else None,
    }
data = {
    "currency": "AUD", "gst": i["B13"].value,
    "plans": {"starter": i["B15"].value, "growth": i["B16"].value, "scale": i["B17"].value},
    "bundleDiscount": i["B10"].value, "setupDiscount": i["B11"].value, "items": items,
    "source": "DUSA-Growth-Engine-Catalogue.xlsx (regenerate with scripts/export-catalogue.py)",
}
with open(out, "w") as f:
    json.dump(data, f, indent=1)
print(f"wrote {len(items)} items to {out}")
